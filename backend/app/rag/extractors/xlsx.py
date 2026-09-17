"""Extractor de XLSX con openpyxl.

Política:
  - Cada hoja se divide en BLOQUES de N filas (ROWS_PER_PAGE).
  - Cada bloque se emite como una ExtractedPage independiente.
  - El header de columnas se repite al inicio de cada bloque para que el
    chunker y el LLM sepan qué significa cada columna.
  - Solo valores con `data_only=True` (valores calculados, no fórmulas).
  - Las hojas vacías se omiten.
  - Cortamos a un máximo de filas por hoja para no explotar con Excels
    gigantes. Configurable.
"""
from __future__ import annotations

from pathlib import Path

from openpyxl import load_workbook

from app.observability.logging import get_logger
from app.rag.extractors.base import ExtractedPage, ExtractionError

log = get_logger(__name__)

# Límite defensivo: XLSX gigantes se truncan. Ajustable.
MAX_ROWS_PER_SHEET = 5_000
MAX_COLS_PER_ROW = 50

# Filas por bloque. Cada bloque se convierte en un ExtractedPage.
# 50 filas × ~250 chars/fila ≈ 12.500 chars → ~5 chunks de 2.400 chars.
ROWS_PER_PAGE = 50


class XlsxExtractor:
    name = "xlsx"
    supported_mimes = (
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        "application/vnd.ms-excel",  # legacy .xls (openpyxl no lo abre; fallará)
    )

    def can_handle(self, mime_type: str) -> bool:
        return mime_type in self.supported_mimes

    def extract(self, path: Path) -> list[ExtractedPage]:
        if not path.is_file():
            raise ExtractionError(f"archivo no encontrado: {path}")
        try:
            wb = load_workbook(str(path), data_only=True, read_only=True)
        except Exception as exc:
            raise ExtractionError(f"XLSX inválido o corrupto: {exc}") from exc

        pages: list[ExtractedPage] = []
        try:
            for sheet_name in wb.sheetnames:
                ws = wb[sheet_name]
                pages.extend(self._extract_sheet(ws, sheet_name))
        finally:
            wb.close()

        return pages

    def _extract_sheet(self, ws, sheet_name: str) -> list[ExtractedPage]:
        """Extrae una hoja como múltiples ExtractedPage de ROWS_PER_PAGE filas."""
        all_rows: list[str] = []
        header: str | None = None
        row_count = 0

        for row in ws.iter_rows(values_only=True):
            if row_count >= MAX_ROWS_PER_SHEET:
                log.info(
                    "xlsx_sheet_truncated",
                    sheet=sheet_name,
                    max_rows=MAX_ROWS_PER_SHEET,
                )
                break

            cells = []
            for cell in row[:MAX_COLS_PER_ROW]:
                if cell is None:
                    cells.append("")
                else:
                    # Normalizar saltos internos para no romper el chunker
                    cells.append(
                        str(cell).replace("\n", " ").replace("\t", " ").strip()
                    )

            # Saltar filas completamente vacías
            if not any(c for c in cells):
                continue

            row_text = "\t".join(cells)

            # La primera fila no vacía se considera el header de columnas.
            if header is None:
                header = row_text
                row_count += 1
                continue

            all_rows.append(row_text)
            row_count += 1

        if not all_rows:
            # Hoja sin datos, solo header o vacía
            if header:
                return [
                    ExtractedPage(
                        text=header,
                        meta={"sheet": sheet_name, "rows": 0, "block": 0},
                    )
                ]
            return []

        # Emitir bloques de ROWS_PER_PAGE, repitiendo el header
        pages: list[ExtractedPage] = []
        total_data_rows = len(all_rows)

        for block_idx, start in enumerate(range(0, total_data_rows, ROWS_PER_PAGE)):
            block = all_rows[start : start + ROWS_PER_PAGE]
            text_lines = [header] if header else []
            text_lines.extend(block)
            text = "\n".join(text_lines)
            pages.append(
                ExtractedPage(
                    text=text,
                    meta={
                        "sheet": sheet_name,
                        "rows": len(block),
                        "block": block_idx,
                        "row_start": start + 1,
                        "row_end": start + len(block),
                    },
                )
            )

        log.info(
            "xlsx_extracted",
            sheet=sheet_name,
            rows=total_data_rows,
            blocks=len(pages),
        )
        return pages